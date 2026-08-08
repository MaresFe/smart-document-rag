import shutil
from pathlib import Path

import fitz
import pandas as pd
import pytest
from docx import Document as DocxDocument
from PIL import Image, ImageDraw, ImageFont

from app.services import text_extraction
from app.services.text_extraction import (
    TextExtractionError,
    extract_text_from_document,
)


def test_txt_extraction_supports_utf8_bom(
    tmp_path: Path,
) -> None:
    path = tmp_path / "belge.txt"
    path.write_text(
        "Belgenin sorumlusu Ayşe Yılmaz'dır.",
        encoding="utf-8-sig",
    )

    result = extract_text_from_document(
        file_type="txt",
        storage_path=str(path),
    )

    assert "Ayşe Yılmaz" in result


def test_csv_extraction_supports_cp1254_and_semicolon(
    tmp_path: Path,
) -> None:
    path = tmp_path / "subeler.csv"
    path.write_bytes(
        (
            "Şube;Şehir;Çalışan Sayısı\n"
            "Merkez;İstanbul;42\n"
            "Doğu;Malatya;18\n"
        ).encode("cp1254"),
    )

    result = extract_text_from_document(
        file_type="csv",
        storage_path=str(path),
    )

    assert "Şube: Doğu" in result
    assert "Şehir: Malatya" in result
    assert "Çalışan Sayısı: 18" in result


def test_xlsx_extraction_includes_multiple_sheets(
    tmp_path: Path,
) -> None:
    path = tmp_path / "rapor.xlsx"

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame(
            {
                "Proje": ["Atlas"],
                "Bütçe": [245000],
            }
        ).to_excel(
            writer,
            sheet_name="Projeler",
            index=False,
        )

        pd.DataFrame(
            {
                "Çalışan": ["Deniz Kaya"],
                "Departman": ["Yazılım"],
            }
        ).to_excel(
            writer,
            sheet_name="Çalışanlar",
            index=False,
        )

    result = extract_text_from_document(
        file_type="xlsx",
        storage_path=str(path),
    )

    assert "[Çalışma Sayfası: Projeler]" in result
    assert "Proje: Atlas" in result
    assert "[Çalışma Sayfası: Çalışanlar]" in result
    assert "Çalışan: Deniz Kaya" in result


def test_docx_extraction_includes_paragraphs_and_tables(
    tmp_path: Path,
) -> None:
    path = tmp_path / "destek.docx"
    document = DocxDocument()
    document.add_paragraph(
        "Destek ekibinin yöneticisi Mert Acar'dır.",
    )

    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Bilgi"
    table.rows[0].cells[1].text = "Değer"

    row = table.add_row().cells
    row[0].text = "Dahili telefon"
    row[1].text = "4407"

    document.save(path)

    result = extract_text_from_document(
        file_type="docx",
        storage_path=str(path),
    )

    assert "Mert Acar" in result
    assert "[Tablo 1]" in result
    assert "Dahili telefon | 4407" in result


def test_pdf_extraction_preserves_page_labels(
    tmp_path: Path,
) -> None:
    path = tmp_path / "proje.pdf"
    document = fitz.open()

    first_page = document.new_page()
    first_page.insert_text(
        (72, 72),
        "Project owner: Deniz Kaya.",
    )

    second_page = document.new_page()
    second_page.insert_text(
        (72, 72),
        "Deadline: 30 October 2026.",
    )

    document.save(path)
    document.close()

    result = extract_text_from_document(
        file_type="pdf",
        storage_path=str(path),
    )

    assert "[Sayfa 1]" in result
    assert "Deniz Kaya" in result
    assert "[Sayfa 2]" in result
    assert "30 October 2026" in result


def test_image_extraction_uses_ocr_result(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "taranmis.png"

    with Image.new("RGB", (800, 400), "white") as image:
        image.save(path, format="PNG")

    monkeypatch.setattr(
        text_extraction.pytesseract,
        "image_to_string",
        lambda *args, **kwargs: (
            "Belgenin sorumlusu Zeynep Aydın'dır."
        ),
    )

    result = extract_text_from_document(
        file_type="png",
        storage_path=str(path),
    )

    assert result.startswith("[Görsel - OCR]")
    assert "Zeynep Aydın" in result


def test_image_without_readable_text_is_rejected(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "metinsiz.jpg"

    with Image.new("RGB", (800, 400), "white") as image:
        image.save(path, format="JPEG")

    monkeypatch.setattr(
        text_extraction.pytesseract,
        "image_to_string",
        lambda *args, **kwargs: "",
    )

    with pytest.raises(TextExtractionError) as error_info:
        extract_text_from_document(
            file_type="jpg",
            storage_path=str(path),
        )

    assert "OCR ile okunabilir metin" in str(
        error_info.value,
    )


@pytest.mark.system_ocr
def test_real_tesseract_reads_numeric_content(
    tmp_path: Path,
) -> None:
    if shutil.which("tesseract") is None:
        pytest.skip("Tesseract sistemde kurulu değil.")

    font_path = Path(
        "/usr/share/fonts/truetype/dejavu/"
        "DejaVuSans.ttf"
    )

    if not font_path.exists():
        pytest.skip("OCR test yazı tipi sistemde bulunamadı.")

    path = tmp_path / "telefon.png"
    font = ImageFont.truetype(str(font_path), 54)

    with Image.new("RGB", (1400, 500), "white") as image:
        draw = ImageDraw.Draw(image)
        draw.text(
            (70, 170),
            "DESTEK TELEFONU 7788",
            fill="black",
            font=font,
        )
        image.save(path, format="PNG")

    result = extract_text_from_document(
        file_type="png",
        storage_path=str(path),
    )

    assert "7788" in result
